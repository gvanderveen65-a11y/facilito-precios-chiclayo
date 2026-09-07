import io
import os
from datetime import datetime
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SOURCE_URL = (
    "https://www.osinergmin.gob.pe/seccion/centro_documental/hidrocarburos/"
    "SCOP/SCOP-DOCS/2026/Registro-precios/Ultimos-Precios-Registrados-EVPC.xlsx"
)
OUT_PATH = "historial/Chiclayo_GasoholRegular_historial.xlsx"
DAILY_PATH = "historial/daily/latest_daily.xlsx"

LIMA = ZoneInfo("America/Lima")
PROVINCIA_OBJETIVO = "CHICLAYO"
PRODUCTO_OBJETIVO = "GASOHOL REGULAR"

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")

# El servidor bloquea clientes sin User-Agent de navegador (regla generica del WAF,
# no es un desafio de comportamiento/captcha) - un User-Agent normal es suficiente.
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
    )
}


def descargar_y_filtrar():
    resp = requests.get(SOURCE_URL, headers=HEADERS, timeout=60)
    resp.raise_for_status()

    wb_src = load_workbook(io.BytesIO(resp.content), read_only=True, data_only=True)
    ws_src = wb_src.active

    header_row = next(ws_src.iter_rows(min_row=1, max_row=1, values_only=True))
    idx = {name: i for i, name in enumerate(header_row)}

    records = []
    for row in ws_src.iter_rows(min_row=2, values_only=True):
        provincia = row[idx["PROVINCIA"]]
        producto = row[idx["PRODUCTO"]]
        if provincia != PROVINCIA_OBJETIVO or producto != PRODUCTO_OBJETIVO:
            continue

        distrito = (row[idx["DISTRITO"]] or "").strip()
        establecimiento = (row[idx["RAZON"]] or "").strip()
        direccion = (row[idx["DIRECCION"]] or "").strip()
        precio = float(row[idx["PRECIO_VENTA"]])

        records.append(
            {
                "Distrito": distrito,
                "Establecimiento": establecimiento,
                "Direccion": direccion,
                "Precio": precio,
                "Key": f"{establecimiento}|{direccion}",
            }
        )

    wb_src.close()
    return records


def style_header_cell(cell):
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")


def update_workbook(records, now):
    sheet_name = now.strftime("%Y-%m")
    col_header = now.strftime("%Y-%m-%d %H:%M")

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)

    if os.path.exists(OUT_PATH):
        wb = load_workbook(OUT_PATH)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
        records_sorted = sorted(records, key=lambda r: r["Precio"])

        headers = ["Distrito", "Establecimiento", "Direccion", col_header]
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            style_header_cell(cell)

        for r, rec in enumerate(records_sorted, start=2):
            ws.cell(row=r, column=1, value=rec["Distrito"]).font = Font(name=FONT_NAME)
            ws.cell(row=r, column=2, value=rec["Establecimiento"]).font = Font(name=FONT_NAME)
            ws.cell(row=r, column=3, value=rec["Direccion"]).font = Font(name=FONT_NAME)
            pc = ws.cell(row=r, column=4, value=rec["Precio"])
            pc.font = Font(name=FONT_NAME)
            pc.number_format = '"S/" #,##0.00'

        widths = [18, 42, 55, 16]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:D{ws.max_row}"
    else:
        ws = wb[sheet_name]
        last_row = ws.max_row
        last_col = ws.max_column
        new_col = last_col + 1

        hcell = ws.cell(row=1, column=new_col, value=col_header)
        style_header_cell(hcell)
        ws.column_dimensions[get_column_letter(new_col)].width = 16

        key_to_row = {}
        for r in range(2, last_row + 1):
            key = f"{ws.cell(row=r, column=2).value}|{ws.cell(row=r, column=3).value}"
            key_to_row[key] = r

        append_row = last_row + 1
        for rec in records:
            if rec["Key"] in key_to_row:
                row = key_to_row[rec["Key"]]
            else:
                row = append_row
                ws.cell(row=row, column=1, value=rec["Distrito"]).font = Font(name=FONT_NAME)
                ws.cell(row=row, column=2, value=rec["Establecimiento"]).font = Font(name=FONT_NAME)
                ws.cell(row=row, column=3, value=rec["Direccion"]).font = Font(name=FONT_NAME)
                key_to_row[rec["Key"]] = row
                append_row += 1

            pc = ws.cell(row=row, column=new_col, value=rec["Precio"])
            pc.font = Font(name=FONT_NAME)
            pc.number_format = '"S/" #,##0.00'

        final_last_row = append_row - 1
        ws.auto_filter.ref = f"A1:{get_column_letter(new_col)}{final_last_row}"

    if "Notas" not in wb.sheetnames:
        wsn = wb.create_sheet("Notas")
        wsn.cell(row=1, column=1, value=(
            "Fuente: Osinergmin - Registro de Precios (PRICE), archivo oficial "
            "Ultimos-Precios-Registrados-EVPC.xlsx (Provincia Chiclayo, Producto Gasohol "
            "Regular). Precios reportados por los propios operadores. Cada hoja mensual "
            "(AAAA-MM) contiene una fila por grifo y una columna por corrida (encabezado "
            "= fecha y hora de la consulta, hora Peru). Actualizado automaticamente via "
            "GitHub Actions cada 3 horas."
        )).font = Font(name=FONT_NAME)
        wsn.column_dimensions["A"].width = 100
        wsn.cell(row=1, column=1).alignment = Alignment(wrap_text=True)

    wb.save(OUT_PATH)
    print(f"Hoja: {sheet_name} | Columna: {col_header} | Grifos: {len(records)}")
    return wb


def build_daily_snapshot(wb, now):
    """Archivo liviano con solo las columnas (corridas) de HOY, recalculado desde
    cero cada vez a partir del historial completo. Se sube a OneDrive cada 3
    horas y se archiva con la fecha del dia una vez pasada la medianoche, sin
    tocar nunca el historial completo (que vive solo en este repo)."""
    sheet_name = now.strftime("%Y-%m")
    today_str = now.strftime("%Y-%m-%d")

    ws_src = wb[sheet_name]
    last_col = ws_src.max_column
    last_row = ws_src.max_row

    today_cols = []
    for c in range(4, last_col + 1):
        header = ws_src.cell(row=1, column=c).value
        if isinstance(header, str) and header.startswith(today_str):
            today_cols.append(c)

    wb_daily = Workbook()
    ws = wb_daily.active
    ws.title = "Evolutivo"

    headers = ["Distrito", "Establecimiento", "Direccion"] + [
        ws_src.cell(row=1, column=c).value for c in today_cols
    ]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        style_header_cell(cell)

    out_r = 2
    for r in range(2, last_row + 1):
        distrito = ws_src.cell(row=r, column=1).value
        establecimiento = ws_src.cell(row=r, column=2).value
        direccion = ws_src.cell(row=r, column=3).value
        precios = [ws_src.cell(row=r, column=c).value for c in today_cols]
        if not any(p is not None for p in precios):
            continue

        ws.cell(row=out_r, column=1, value=distrito)
        ws.cell(row=out_r, column=2, value=establecimiento)
        ws.cell(row=out_r, column=3, value=direccion)
        for i, p in enumerate(precios):
            pc = ws.cell(row=out_r, column=4 + i, value=p)
            pc.number_format = '"S/" #,##0.00'
        out_r += 1

    final_last_row = out_r - 1
    widths = [18, 42, 55] + [16] * len(today_cols)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    if final_last_row >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(3 + len(today_cols))}{final_last_row}"

    os.makedirs(os.path.dirname(DAILY_PATH), exist_ok=True)
    wb_daily.save(DAILY_PATH)
    print(f"Snapshot diario: {today_str} | Columnas de hoy: {len(today_cols)} | Filas: {final_last_row}")


if __name__ == "__main__":
    now = datetime.now(LIMA)
    records = descargar_y_filtrar()
    if len(records) < 50:
        raise SystemExit(f"Muy pocos registros ({len(records)}) - revisar el archivo fuente, no se guarda.")
    wb = update_workbook(records, now)
    build_daily_snapshot(wb, now)
